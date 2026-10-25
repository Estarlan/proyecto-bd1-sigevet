# Created: 23:50 08-Oct-2026
# Updated: 01:05 09-Oct-2026
# Updated: 22:55 09-Oct-2026
# Updated: 13:03 10-Oct-2026
# Updated: 13:25 10-Oct-2026
# Updated: 18:55 24-Oct-2026
# Updated: 18:55 24-Oct-2026
"""Convierte docs/src/**/*.md en documentos Word con portada institucional.

Uso:
    python docs/tools/build_docs.py                 # todos los documentos
    python docs/tools/build_docs.py docs/src/fase1/01_Propuesta.md
Requisitos: pandoc en el PATH, python-docx, pyyaml. La plantilla de estilos se crea con make_reference.py.
"""
import subprocess
import sys
from pathlib import Path

import yaml
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

REPO = Path(__file__).resolve().parents[2]
SRC = REPO / "docs" / "src"
REFERENCE = REPO / "docs" / "tools" / "reference.docx"
LOGO = REPO / "docs" / "assets" / "logo_unbosque.png"
FIRMAS = SRC / "firmas"   # juan_david.png, sebastian.png, juan_miguel.png, docente.png (opcionales)
ARCHIVO_FIRMA = {"Juan David": "juan_david", "Sebastián": "sebastian", "Juan Miguel": "juan_miguel", "Christian": "docente"}
CARPETA_SALIDA = {"fase1": "Fase1", "fase2": "Fase2", "transversales": "Transversales", "final": "Final"}
NARANJA = RGBColor(0xC8, 0x6A, 0x00)
GRIS = RGBColor(0x33, 0x33, 0x33)
ESTILOS_TITULO = {"Title", "Subtitle", "Author", "Date", "Abstract"}
PANDOC_FORMATO = "markdown+pipe_tables+grid_tables+table_captions+implicit_figures+tex_math_dollars+subscript+superscript"


def leer_front_matter(md: Path) -> dict:
    texto = md.read_text(encoding="utf-8")
    if not texto.startswith("---"):
        raise ValueError(f"{md} no tiene encabezado YAML")
    return yaml.safe_load(texto.split("---", 2)[1]) or {}


def expandir_inclusiones(md: Path) -> Path:
    """Reemplaza líneas `<!-- INCLUDE: ruta -->` (relativa a docs/src) por el contenido del archivo."""
    lineas = []
    for linea in md.read_text(encoding="utf-8").splitlines():
        marca = linea.strip()
        if marca.startswith("<!-- INCLUDE:") and marca.endswith("-->"):
            ruta = SRC / marca[len("<!-- INCLUDE:"):-3].strip()
            lineas.append(ruta.read_text(encoding="utf-8"))
        else:
            lineas.append(linea)
    expandido = md.with_name(f".{md.stem}.expandido.md")
    expandido.write_text("\n".join(lineas) + "\n", encoding="utf-8")
    return expandido


def ejecutar_pandoc(md: Path, salida: Path) -> None:
    fuente = expandir_inclusiones(md)
    cmd = [
        "pandoc", str(fuente), "-f", PANDOC_FORMATO, "-o", str(salida),
        "--reference-doc", str(REFERENCE), "--resource-path", f"{REPO}:{md.parent}",
        "--toc", "--toc-depth=2", "--number-sections",
        "-M", "toc-title=Tabla de contenido", "-M", "lang=es-CO",
    ]
    try:
        subprocess.run(cmd, check=True, cwd=REPO)
    finally:
        fuente.unlink(missing_ok=True)


def _parrafo(doc, texto="", tam=12, negrita=False, color=GRIS, espacio_despues=6):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(espacio_despues)
    if texto:
        r = p.add_run(texto)
        r.font.size = Pt(tam)
        r.font.bold = negrita
        r.font.color.rgb = color
    return p


def construir_portada(doc, meta: dict, portada: dict) -> list:
    """Crea los párrafos de portada al final del documento y los devuelve para moverlos al inicio."""
    fase = meta.get("fase", "Transversal")
    fecha = meta.get("fecha") or portada.get("fechas", {}).get(fase, "")   # sin fecha si no se define
    elementos = []
    logo = _parrafo(doc, espacio_despues=18)
    logo.add_run().add_picture(str(LOGO), width=Cm(6))
    elementos.append(logo)
    elementos.append(_parrafo(doc, portada["universidad"], 16, True, NARANJA, 2))
    elementos.append(_parrafo(doc, portada["facultad"], 11))
    elementos.append(_parrafo(doc, portada["asignatura"], 11, espacio_despues=60))
    elementos.append(_parrafo(doc, meta.get("subtitle", ""), 22, True, NARANJA, 10))
    elementos.append(_parrafo(doc, portada["proyecto"], 14, True))
    elementos.append(_parrafo(doc, portada["organizacion"], 11))
    etiqueta = f"Proyecto Final · {fase}" + (f" · Versión {meta['version']}" if meta.get("version") else "")
    elementos.append(_parrafo(doc, etiqueta, 11, espacio_despues=60))
    elementos.append(_parrafo(doc, "Integrantes", 12, True, espacio_despues=2))
    for nombre in portada["integrantes"]:
        elementos.append(_parrafo(doc, nombre, 12, espacio_despues=2))
    elementos.append(_parrafo(doc, "", espacio_despues=24))
    elementos.append(_parrafo(doc, "Docente", 12, True, espacio_despues=2))
    elementos.append(_parrafo(doc, portada["docente"], 12, espacio_despues=48))
    cierre = _parrafo(doc, portada["ciudad"] + (f", {fecha}" if fecha else ""), 12)
    cierre.add_run().add_break(WD_BREAK.PAGE)
    elementos.append(cierre)
    return elementos


def formatear_tablas(doc) -> None:
    for tabla in doc.tables:
        tbl_pr = tabla._tbl.tblPr
        bordes = OxmlElement("w:tblBorders")
        for lado in ("top", "left", "bottom", "right", "insideH", "insideV"):
            b = OxmlElement(f"w:{lado}")
            b.set(qn("w:val"), "single")
            b.set(qn("w:sz"), "4")
            b.set(qn("w:color"), "A6A6A6")
            bordes.append(b)
        tbl_pr.append(bordes)
        if not tabla.rows:
            continue
        for celda in tabla.rows[0].cells:
            sombra = OxmlElement("w:shd")
            sombra.set(qn("w:val"), "clear")
            sombra.set(qn("w:fill"), "FBE5CC")
            celda._tc.get_or_add_tcPr().append(sombra)
            for p in celda.paragraphs:
                for r in p.runs:
                    r.font.bold = True
        for fila in tabla.rows:
            for celda in fila.cells:
                for p in celda.paragraphs:
                    for r in p.runs:
                        r.font.size = Pt(9.5)


def insertar_firmas(doc) -> None:
    """En las tablas con columna Firma: deja espacio para firmar o pega la imagen si existe en docs/src/firmas/."""
    for tabla in doc.tables:
        encabezado = [c.text.strip() for c in tabla.rows[0].cells]
        if "Firma" not in encabezado:
            continue
        col = encabezado.index("Firma")
        for fila in tabla.rows[1:]:
            fila.height = Cm(1.6)
            nombre = fila.cells[0].text
            clave = next((v for k, v in ARCHIVO_FIRMA.items() if k in nombre), None)
            imagen = FIRMAS / f"{clave}.png" if clave else None
            if imagen and imagen.exists():
                celda = fila.cells[col]
                celda.paragraphs[0].add_run().add_picture(str(imagen), height=Cm(1.3))


def postprocesar(docx: Path, meta: dict, portada: dict) -> None:
    doc = Document(str(docx))
    cuerpo = doc.element.body
    for p in list(doc.paragraphs):
        if p.style is not None and p.style.name in ESTILOS_TITULO:
            cuerpo.remove(p._p)
    portada_elems = construir_portada(doc, meta, portada)
    for i, p in enumerate(portada_elems):
        cuerpo.insert(i, p._p)
    toc = cuerpo.find(qn("w:sdt"))
    if toc is not None:
        salto = doc.add_paragraph()
        salto.add_run().add_break(WD_BREAK.PAGE)
        toc.addnext(salto._p)
    formatear_tablas(doc)
    insertar_firmas(doc)
    actualizar = OxmlElement("w:updateFields")   # Word actualiza la tabla de contenido al abrir el archivo
    actualizar.set(qn("w:val"), "true")
    doc.settings.element.append(actualizar)
    doc.core_properties.title = f"{meta.get('subtitle', '')} — SIGEVET"
    doc.core_properties.author = ", ".join(n.split(" — ")[0] for n in portada["integrantes"])
    # Keep tall Draw.io sheets inside the printable area, preserving their ratio.
    section = doc.sections[0]
    max_width = (section.page_width or Cm(21.59)) - (section.left_margin or Cm(2.54)) - (section.right_margin or Cm(2.54))
    max_height = (section.page_height or Cm(27.94)) - (section.top_margin or Cm(2.54)) - (section.bottom_margin or Cm(2.54)) - Cm(2)
    for picture in doc.inline_shapes:
        if picture.width and picture.height:
            factor = min(1, max_width / picture.width, max_height / picture.height)
            if factor < 1:
                picture.width, picture.height = int(picture.width * factor), int(picture.height * factor)
    history = doc.core_properties.comments or "Created: [unknown]"
    doc.core_properties.comments = history + "\nUpdated: 18:55 24-Oct-2026 — modelos Draw.io"
    doc.save(str(docx))


def construir(md: Path, portada: dict) -> Path:
    meta = leer_front_matter(md)
    carpeta = REPO / "entregables" / CARPETA_SALIDA[md.parent.name]
    carpeta.mkdir(parents=True, exist_ok=True)
    salida = carpeta / f"{md.stem}.docx"
    ejecutar_pandoc(md, salida)
    postprocesar(salida, meta, portada)
    return salida


def main(argv: list) -> int:
    portada = yaml.safe_load((SRC / "portada.yaml").read_text(encoding="utf-8"))
    archivos = [Path(a).resolve() for a in argv] or sorted(SRC.glob("fase*/[0-9]*.md")) + sorted(SRC.glob("transversales/[0-9]*.md"))
    errores = 0
    for md in archivos:
        try:
            print(f"OK  {construir(md, portada).relative_to(REPO)}")
        except (subprocess.CalledProcessError, ValueError, KeyError) as exc:
            errores += 1
            print(f"ERR {md.relative_to(REPO)}: {exc}")
    print(f"Documentos generados: {len(archivos) - errores}/{len(archivos)}")
    return 1 if errores else 0


if __name__ == "__main__":
    sys.exit(main(sys.argv[1:]))
